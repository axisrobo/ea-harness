"""
from_document.py — Extract architecture requirements from text documents.

Supports: .pdf  .docx  .md  .txt

Strategy: extract text, then use the Claude API to identify architecture-relevant
information and map it onto the req/v2 entity model (infra / systems / components
/ deployments / flows / network_links / auth). References between entities are
emitted as names; the merger resolves them to typed IDs.

Usage:
    python from_document.py -i requirements.pdf -o partial-req.yaml
    python from_document.py -i design-doc.docx -o partial-req.yaml
    ANTHROPIC_API_KEY=... python from_document.py -i brief.txt -o partial-req.yaml
"""

import argparse
import json
import os
import re
import sys
from pathlib import Path

from .normalizer import (
    Confidence, PartialReq, partial_req_to_yaml, v2_json_to_partial,
)


# ── Text extraction ───────────────────────────────────────────────────────────

def _extract_text_pdf(path: str) -> str:
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            pages = [p.extract_text() or "" for p in pdf.pages]
        return "\n".join(pages)
    except ImportError:
        try:
            from pypdf import PdfReader
            reader = PdfReader(path)
            return "\n".join(p.extract_text() or "" for p in reader.pages)
        except ImportError:
            return ""


def _extract_text_docx(path: str) -> str:
    try:
        from docx import Document
        doc = Document(path)
        parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except ImportError:
        return ""


def extract_text(path: str) -> str:
    """Extract plain text from a document file."""
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        text = _extract_text_pdf(path)
    elif ext in (".docx", ".doc"):
        text = _extract_text_docx(path)
    else:
        with open(path, encoding="utf-8", errors="replace") as f:
            text = f.read()
    return text


# ── LLM extraction ────────────────────────────────────────────────────────────

EXTRACTION_PROMPT = """You are an enterprise architecture analyst extracting structured requirements from a document.

Read the document and extract every architecture-relevant fact into the JSON shape below.
Physical precision matters more than completeness: use null for anything not explicitly stated.

Critical modelling rules:
- `infra` holds hosting locations and network topology. node_kind is the topology role
  (region | data_center | iaas_vpc_vnet | paas | saas | third_party | office_network |
  factory_network | lab | internet_network | network_zone | subnet | firewall | waf |
  security_gateway | router | switch | vpn_gateway | identity_provider | soc_monitoring |
  load_balancer | bastion_host | logging_service | policy_service | key_management).
  infra_type is the hosting category (private_cloud | public_cloud | saas | third_party |
  office | factory | lab). network_type is the network/security domain (prod_network | dmz |
  office_network | factory_network | lab_network). Use "parent" to name the containing node.
- Firewalls, WAFs, routers, VPN gateways, load-balancer appliances, bastion hosts and
  identity providers (ADFS, Entra ID) are infra nodes, NEVER components.
- `components` are application services/databases/buses/API gateways only. Use "system"
  to name the owning system, and set component_role to the closest value
  (web_frontend | backend_service | bff | api_gateway | message_bus | database | cache |
  object_storage | integration_service | data_lake | data_warehouse | batch_processing |
  streaming_processing | ...).
- `flows` are directed component-to-component communications (caller -> provider). Use the
  literal "internet" as the source for external ingress. auth_method is service-to-service:
  OAuth2_ClientCredentials | mTLS | ClientCertificate | SASL_SCRAM | Basic | ApiKey |
  UserPassword | Kerberos | IAM_Role | ManagedIdentity | none. Put appliances the path
  traverses in "via" (names).
- `network_links` are undirected infra-to-infra connections. A carrier circuit
  (ExpressRoute / MPLS / Direct Connect) is a link, not a node.
- `auth` is user/entry authentication only (an identity-provider redirect is an auth row,
  not a flow). subject is "user" or "application".
- Encryption: at-rest goes on the component (encryption_at_rest), in-transit goes on the
  flow (encryption = TLS1.3 | TLS1.2 | mTLS | IPSec | none | TBD).
- Every reference between entities is a NAME, not an ID. Do not invent IDs.

Document content:
---
{document_text}
---

Output ONLY this JSON, no explanation. Use null for unknown fields; be conservative and
extract only what the document states.

{{
  "project": {{
    "name": null, "id": null, "scope": "standalone|modification|e2e",
    "department": null, "data_classification": null
  }},
  "infra": [
    {{ "name": null, "node_kind": null, "infra_type": null, "network_type": null,
       "parent": null, "country": null, "vendor": null, "infra_owner": null }}
  ],
  "systems": [
    {{ "name": null, "type": "new|existing|modified", "owner": "org_it|biz_owned|third_party",
       "vendor": null, "data_classification": null }}
  ],
  "components": [
    {{ "system": null, "name": null, "kind": "service|component", "layer": null,
       "component_role": null, "function_desc": null, "sensitivity": null,
       "encryption_at_rest": null, "key_management": null }}
  ],
  "stacks": [
    {{ "component": null, "component_name": null, "component_package": null,
       "version": null, "category": null, "license": null, "eol_date": null,
       "standard_flag": null }}
  ],
  "deployments": [
    {{ "component": null, "environment": "dev|test|staging|prod|dr",
       "deployment_type": "private_cloud|public_cloud|public_cloud_paas|saas|third_party",
       "location_type": "data_center|public_cloud_region|saas", "infra": null,
       "runtime_type": "vm|container|physical|serverless", "runtime_detail": null,
       "instance_count": null }}
  ],
  "flows": [
    {{ "from": null, "to": null, "protocol": null, "port": null, "auth_method": null,
       "encryption": null, "cross_border": null, "via": [], "notes": null }}
  ],
  "network_links": [
    {{ "from": null, "to": null, "method": null, "bandwidth": null, "encrypted": null,
       "encryption_method": null, "managed_by": null, "redundancy": null, "notes": null }}
  ],
  "auth": [
    {{ "subject": "user|application", "applies_to": null, "auth_server": null,
       "protocol": "OIDC|OAuth2_AuthCode|SAML2|CAS|Kerberos|Basic|ApiKey",
       "authorization": "RBAC|ABAC|PBAC|DAC", "authorization_platform": null,
       "user_roles": [], "mfa": null, "notes": null }}
  ],
  "ecosystem_relations": [
    {{ "from": null, "to": null, "relation_type": "upstream|downstream|partner|customer",
       "notes": null }}
  ],
  "credentials": [ {{ "environment": "azure|aws|private_dc|saas|other", "solution": null, "notes": null }} ],
  "constraints": [],
  "open_items": [ {{ "id": null, "description": null, "owner": null, "blocking": null }} ],
  "gaps_noted": ["information that seems important but is missing from the document"]
}}
"""


def call_llm_extraction(document_text: str, source: str) -> dict:
    """Call the Claude API to extract structured requirements from document text."""
    api_key = os.environ.get("ANTHROPIC_API_KEY", "")
    if not api_key:
        return {"error": "ANTHROPIC_API_KEY not set"}

    import urllib.request

    max_chars = 80000
    if len(document_text) > max_chars:
        document_text = document_text[:max_chars] + "\n... [document truncated]"

    prompt = EXTRACTION_PROMPT.format(document_text=document_text)

    payload = json.dumps({
        "model": "claude-sonnet-4-6",
        "max_tokens": 8192,
        "messages": [{"role": "user", "content": prompt}]
    }).encode("utf-8")

    req = urllib.request.Request(
        "https://api.anthropic.com/v1/messages",
        data=payload,
        headers={
            "x-api-key": api_key,
            "anthropic-version": "2023-06-01",
            "content-type": "application/json",
        },
        method="POST"
    )

    try:
        with urllib.request.urlopen(req, timeout=90) as resp:
            response = json.loads(resp.read().decode("utf-8"))
    except Exception as e:
        return {"error": str(e)}

    text = ""
    for block in response.get("content", []):
        if block.get("type") == "text":
            text = block["text"]
            break

    json_match = re.search(r'\{[\s\S]+\}', text)
    if not json_match:
        return {"error": "No JSON found in LLM response"}

    try:
        return json.loads(json_match.group())
    except json.JSONDecodeError as e:
        return {"error": f"JSON parse error: {e}"}


# ── Map LLM output → PartialReq ───────────────────────────────────────────────

def map_to_partial_req(extracted: dict, source_file: str) -> PartialReq:
    req = v2_json_to_partial(
        extracted,
        source_tool="arch-req-from-doc",
        source_file=source_file,
        confidence=Confidence.MEDIUM,
        source_label=f"document:{Path(source_file).name}",
    )
    req.gaps.append(
        "Document extraction is MEDIUM confidence — verify all values, especially "
        "authentication mechanisms and physical locations"
    )
    return req


# ── Main ─────────────────────────────────────────────────────────────────────

def parse_document(input_path: str) -> PartialReq:
    text = extract_text(input_path)
    if not text.strip():
        req = PartialReq(source_tool="arch-req-from-doc", source_file=input_path)
        req.gaps.append(f"Could not extract text from {input_path}. "
                        f"Install pdfplumber (PDF) or python-docx (DOCX).")
        return req
    extracted = call_llm_extraction(text, input_path)
    return map_to_partial_req(extracted, input_path)


def main():
    parser = argparse.ArgumentParser(description="Extract requirements from document")
    parser.add_argument("-i", "--input", required=True, help="Input document file")
    parser.add_argument("-o", "--output", default=None, help="Output partial-req YAML")
    args = parser.parse_args()

    req = parse_document(args.input)
    out = partial_req_to_yaml(req)

    if args.output:
        with open(args.output, "w", encoding="utf-8") as f:
            f.write(out)
        print(f"✓ Partial requirements written: {args.output}")
        for label, items in (
            ("Infra nodes", req.infra), ("Systems", req.systems),
            ("Components", req.components), ("Deployments", req.deployments),
            ("Flows", req.flows), ("Network links", req.network_links),
            ("Auth entries", req.auth),
        ):
            print(f"  {label}: {len(items)}")
        if req.gaps:
            print("  Notes:")
            for g in req.gaps[:5]:
                print(f"    • {g}")
    else:
        print(out)


if __name__ == "__main__":
    main()
